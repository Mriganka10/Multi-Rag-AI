from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

from app.artifacts.service import ArtifactExportService


NOTICE_RESPONSE = """SCN Analysis - Income Tax Assessment Year 2018-19

Taxpayer: Raj Shiyaram Verma
PAN: ACVPV2752M
Notice Date: 19/03/2024
DIN: ITBA/AST/F/147(SCN)/2023-24/1062953559(1)
Proceeding: Reassessment proceedings pursuant to notice under section 148
Main Issue: Proposed variation under section 50C relating to sale of immovable property

Summary

The Income Tax Department has issued a Show Cause Notice for professional review.

Key Observations

1. The stamp duty value should be reconciled with the sale consideration.
"""


def test_exporters_preserve_response_title_and_header_fields(tmp_path, monkeypatch) -> None:
    from app.core.config import settings

    monkeypatch.setattr(settings, "data_dir", tmp_path)
    artifacts = ArtifactExportService().generate(
        content=NOTICE_RESPONSE,
        query="Provide all formats",
        agent_name="scn",
        tenant_id="header-test",
        explicit_formats=["pdf", "docx", "xlsx"],
    )
    paths = {artifact.format: artifact.location for artifact in artifacts}

    pdf_text = "\n".join(
        page.extract_text() or "" for page in PdfReader(paths["pdf"]).pages
    )
    assert "SCN Analysis - Income Tax Assessment Year 2018-19" in pdf_text
    assert "Taxpayer: Raj Shiyaram Verma" in pdf_text
    assert "PAN: ACVPV2752M" in pdf_text
    assert "Draft Notice Analysis and Response" not in pdf_text

    document = Document(paths["docx"])
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    assert paragraphs[0] == "SCN Analysis - Income Tax Assessment Year 2018-19"
    assert "Taxpayer: Raj Shiyaram Verma" in paragraphs
    assert "PAN: ACVPV2752M" in paragraphs

    workbook = load_workbook(paths["xlsx"], read_only=True)
    summary = workbook["Executive Summary"]
    assert summary["A1"].value == "SCN Analysis - Income Tax Assessment Year 2018-19"
    assert summary["A5"].value == "Taxpayer"
    assert summary["B5"].value == "Raj Shiyaram Verma"
    assert summary["A6"].value == "PAN"
    assert summary["B6"].value == "ACVPV2752M"
