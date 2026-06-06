import re
import csv
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from PIL import Image
from pypdf import PdfReader

from app.models.schemas import AgentName, TaskResult


class OCRAgent:
    def extract_text(self, file_path: Path) -> str:
        suffix = file_path.suffix.lower()
        if suffix == ".pdf":
            return self._extract_pdf_text(file_path)
        if suffix in {".txt", ".csv"}:
            if suffix == ".csv":
                return self._extract_csv_text(file_path)
            return file_path.read_text(encoding="utf-8", errors="ignore")
        if suffix == ".xlsx":
            return self._extract_xlsx_text(file_path)
        if suffix in {".png", ".jpg", ".jpeg", ".tiff", ".bmp"}:
            return self._extract_image_text(file_path)
        raise ValueError(f"Unsupported file type: {suffix}")

    def run(self, query: str, text: str) -> TaskResult:
        rows = self.parse_transactions(text)
        summary = f"Extracted {len(rows)} transaction-like rows from the document."
        return TaskResult(
            agent=AgentName.OCR,
            summary=summary,
            data={"rows": rows, "query": query},
            requires_human_review=True,
        )

    def export_excel(self, rows: list[dict[str, object]], output_path: Path) -> Path:
        output_path.parent.mkdir(parents=True, exist_ok=True)
        pd.DataFrame(rows).to_excel(output_path, index=False)
        return output_path

    def parse_transactions(self, text: str) -> list[dict[str, object]]:
        rows: list[dict[str, object]] = []
        pattern = re.compile(
            r"(?P<date>\d{4}-\d{2}-\d{2}|\d{2}[/-]\d{2}[/-]\d{4})\s+"
            r"(?P<description>.*?)\s+"
            r"(?P<debit>-?\d+(?:,\d{3})*(?:\.\d+)?)\s+"
            r"(?P<credit>-?\d+(?:,\d{3})*(?:\.\d+)?)\s+"
            r"(?P<balance>-?\d+(?:,\d{3})*(?:\.\d+)?)"
        )
        for line in text.splitlines():
            match = pattern.search(line.strip())
            if not match:
                continue
            row = match.groupdict()
            rows.append(
                {
                    "date": row["date"],
                    "description": row["description"].strip(),
                    "debit": self._to_float(row["debit"]),
                    "credit": self._to_float(row["credit"]),
                    "balance": self._to_float(row["balance"]),
                }
            )
        return rows

    def _extract_pdf_text(self, file_path: Path) -> str:
        reader = PdfReader(str(file_path))
        return "\n".join(page.extract_text() or "" for page in reader.pages)

    def _extract_csv_text(self, file_path: Path) -> str:
        with file_path.open("r", encoding="utf-8", errors="ignore", newline="") as handle:
            return "\n".join(" ".join(cell.strip() for cell in row) for row in csv.reader(handle))

    def _extract_xlsx_text(self, file_path: Path) -> str:
        workbook = load_workbook(file_path, read_only=True, data_only=True)
        lines: list[str] = []
        for worksheet in workbook.worksheets:
            for row in worksheet.iter_rows(values_only=True):
                values = [self._format_cell_value(value) for value in row if value is not None]
                if values:
                    lines.append(" ".join(values))
        return "\n".join(lines)

    def _extract_image_text(self, file_path: Path) -> str:
        try:
            import pytesseract
        except ImportError as exc:
            raise RuntimeError("pytesseract is required for local image OCR") from exc
        return pytesseract.image_to_string(Image.open(file_path))

    @staticmethod
    def _to_float(value: str) -> float:
        return float(value.replace(",", ""))

    @staticmethod
    def _format_cell_value(value: object) -> str:
        if isinstance(value, float) and value.is_integer():
            return str(int(value))
        return str(value)
