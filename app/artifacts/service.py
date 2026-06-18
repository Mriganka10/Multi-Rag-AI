import re
from dataclasses import dataclass
from datetime import UTC, datetime
from html import escape
from pathlib import Path
from typing import Iterable
from uuid import uuid4

from app.core.audit import safe_tenant_id
from app.core.config import settings
from app.core.storage import UploadStorage


SUPPORTED_FORMATS = ("pdf", "docx", "xlsx")
_LOCAL_ARTIFACTS: dict[str, "GeneratedArtifact"] = {}


@dataclass(frozen=True)
class GeneratedArtifact:
    artifact_id: str
    tenant_id: str
    format: str
    filename: str
    content_type: str
    location: str

    def public_payload(self) -> dict[str, str]:
        return {
            "id": self.artifact_id,
            "format": self.format,
            "label": {
                "pdf": "Download PDF",
                "docx": "Download Word",
                "xlsx": "Download Excel",
            }[self.format],
            "filename": self.filename,
            "url": f"/api/v1/artifacts/{self.artifact_id}/download",
        }


class ArtifactExportService:
    def __init__(self) -> None:
        self.storage = UploadStorage()

    def requested_formats(
        self,
        query: str,
        explicit_formats: Iterable[str] | None = None,
    ) -> list[str]:
        requested = {
            item.lower().strip()
            for item in (explicit_formats or [])
            if item.lower().strip() in SUPPORTED_FORMATS
        }
        normalized = query.lower()
        if re.search(r"\bpdf\b", normalized):
            requested.add("pdf")
        if re.search(r"\b(word|docx)\b", normalized):
            requested.add("docx")
        if re.search(r"\b(excel|xlsx|spreadsheet)\b", normalized):
            requested.add("xlsx")
        if re.search(r"\b(all formats|all three|pdf.*word.*excel|pdf.*excel.*word)\b", normalized):
            requested.update(SUPPORTED_FORMATS)
        return [item for item in SUPPORTED_FORMATS if item in requested]

    def generate(
        self,
        *,
        content: str,
        query: str,
        agent_name: str,
        tenant_id: str,
        explicit_formats: Iterable[str] | None = None,
    ) -> list[GeneratedArtifact]:
        formats = self.requested_formats(query, explicit_formats)
        if not formats:
            return []

        title = _document_title(agent_name)
        artifacts: list[GeneratedArtifact] = []
        for output_format in formats:
            artifact_id = uuid4().hex
            output_dir = (
                settings.data_dir
                / "outputs"
                / safe_tenant_id(tenant_id)
                / artifact_id
            )
            output_dir.mkdir(parents=True, exist_ok=True)
            filename = f"{_slug(title)}-{artifact_id[:8]}.{output_format}"
            local_path = output_dir / filename
            if output_format == "pdf":
                _write_pdf(local_path, title, content)
            elif output_format == "docx":
                _write_docx(local_path, title, content)
            else:
                _write_xlsx(local_path, title, agent_name, content)

            location = self.storage.persist_artifact(
                local_path=local_path,
                tenant_id=tenant_id,
                artifact_id=artifact_id,
            )
            artifact = GeneratedArtifact(
                artifact_id=artifact_id,
                tenant_id=safe_tenant_id(tenant_id),
                format=output_format,
                filename=filename,
                content_type=_content_type(output_format),
                location=location,
            )
            self._register(artifact)
            artifacts.append(artifact)
        return artifacts

    def get_for_tenant(self, artifact_id: str, tenant_id: str) -> GeneratedArtifact | None:
        tenant = safe_tenant_id(tenant_id)
        if settings.database_url:
            rows = self._fetch_db(
                """
                select artifact_id, tenant_id, format, filename, content_type, location
                from generated_artifacts
                where artifact_id = %s and tenant_id = %s
                limit 1
                """,
                (artifact_id, tenant),
            )
            if not rows:
                return None
            return GeneratedArtifact(*rows[0])
        artifact = _LOCAL_ARTIFACTS.get(artifact_id)
        return artifact if artifact and artifact.tenant_id == tenant else None

    def _register(self, artifact: GeneratedArtifact) -> None:
        if settings.database_url:
            self._with_db(
                """
                insert into generated_artifacts (
                    artifact_id, tenant_id, format, filename, content_type, location, created_at
                )
                values (%s, %s, %s, %s, %s, %s, now())
                """,
                (
                    artifact.artifact_id,
                    artifact.tenant_id,
                    artifact.format,
                    artifact.filename,
                    artifact.content_type,
                    artifact.location,
                ),
            )
            return
        _LOCAL_ARTIFACTS[artifact.artifact_id] = artifact

    def _ensure_table(self) -> None:
        try:
            import psycopg
        except ImportError as exc:
            raise RuntimeError("Install cloud dependencies with: pip install -e '.[cloud]'") from exc
        with psycopg.connect(settings.database_url) as conn:
            with conn.cursor() as cur:
                cur.execute(
                    """
                    create table if not exists generated_artifacts (
                        artifact_id text primary key,
                        tenant_id text not null,
                        format text not null,
                        filename text not null,
                        content_type text not null,
                        location text not null,
                        created_at timestamptz not null
                    )
                    """
                )
            conn.commit()

    def _with_db(self, query: str, params: tuple[object, ...]) -> None:
        self._ensure_table()
        import psycopg

        with psycopg.connect(settings.database_url) as conn:
            with conn.cursor() as cur:
                cur.execute(query, params)
            conn.commit()

    def _fetch_db(self, query: str, params: tuple[object, ...]) -> list[tuple[str, ...]]:
        self._ensure_table()
        import psycopg

        with psycopg.connect(settings.database_url) as conn:
            with conn.cursor() as cur:
                cur.execute(query, params)
                return cur.fetchall()


def _write_pdf(path: Path, title: str, content: str) -> None:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "CAExportTitle",
        parent=styles["Title"],
        fontName="Helvetica-Bold",
        fontSize=18,
        leading=22,
        alignment=TA_CENTER,
        textColor=colors.HexColor("#0F4F45"),
        spaceAfter=14,
    )
    heading_style = ParagraphStyle(
        "CAExportHeading",
        parent=styles["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=12,
        leading=15,
        textColor=colors.HexColor("#176B5D"),
        spaceBefore=9,
        spaceAfter=5,
    )
    body_style = ParagraphStyle(
        "CAExportBody",
        parent=styles["BodyText"],
        fontName="Helvetica",
        fontSize=9.5,
        leading=14,
        textColor=colors.HexColor("#18212A"),
        spaceAfter=5,
    )
    bullet_style = ParagraphStyle(
        "CAExportBullet",
        parent=body_style,
        leftIndent=12,
        firstLineIndent=-7,
    )

    document = SimpleDocTemplate(
        str(path),
        pagesize=A4,
        rightMargin=18 * mm,
        leftMargin=18 * mm,
        topMargin=18 * mm,
        bottomMargin=18 * mm,
        title=title,
        author=settings.app_name,
    )
    story = [Paragraph(escape(title), title_style), Spacer(1, 4)]
    for kind, value in _content_blocks(content):
        safe_value = escape(_pdf_safe(value)).replace("\n", "<br/>")
        if kind == "heading":
            story.append(Paragraph(safe_value, heading_style))
        elif kind == "bullet":
            story.append(Paragraph(f"- {safe_value}", bullet_style))
        elif kind == "numbered":
            story.append(Paragraph(safe_value, bullet_style))
        elif kind == "page_break":
            story.append(PageBreak())
        else:
            story.append(Paragraph(safe_value, body_style))
    document.build(story)


def _write_docx(path: Path, title: str, content: str) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt, RGBColor

    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    title_paragraph = document.add_paragraph()
    title_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_paragraph.add_run(title)
    title_run.bold = True
    title_run.font.name = "Aptos Display"
    title_run.font.size = Pt(18)
    title_run.font.color.rgb = RGBColor(15, 79, 69)

    for kind, value in _content_blocks(content):
        if kind == "heading":
            paragraph = document.add_paragraph()
            run = paragraph.add_run(value)
            run.bold = True
            run.font.name = "Aptos"
            run.font.size = Pt(12)
            run.font.color.rgb = RGBColor(23, 107, 93)
        elif kind == "bullet":
            paragraph = document.add_paragraph(style="List Bullet")
            paragraph.add_run(value)
        elif kind == "numbered":
            paragraph = document.add_paragraph(style="List Number")
            paragraph.add_run(re.sub(r"^\d+\.\s+", "", value))
        elif kind == "page_break":
            document.add_page_break()
        else:
            paragraph = document.add_paragraph()
            paragraph.add_run(value)
        paragraph.paragraph_format.space_after = Pt(5)
        paragraph.paragraph_format.line_spacing = 1.1

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run(
        f"Generated by {settings.app_name} on {datetime.now(UTC).strftime('%d %b %Y %H:%M UTC')}"
    )
    footer_run.font.size = Pt(8)
    footer_run.font.color.rgb = RGBColor(100, 113, 109)
    document.save(path)


def _write_xlsx(path: Path, title: str, agent_name: str, content: str) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    workbook = Workbook()
    summary = workbook.active
    summary.title = "Executive Summary"
    summary.sheet_view.showGridLines = False
    summary["A1"] = title
    summary["A1"].font = Font(size=18, bold=True, color="0F4F45")
    summary["A3"] = "Agent"
    summary["B3"] = agent_name
    summary["A4"] = "Generated"
    summary["B4"] = datetime.now(UTC).strftime("%d %b %Y %H:%M UTC")
    summary["A6"] = "Client-ready analysis"
    summary["A6"].font = Font(size=12, bold=True, color="176B5D")
    summary["A7"] = content
    summary["A7"].alignment = Alignment(wrap_text=True, vertical="top")
    summary.merge_cells("A7:H40")
    summary.row_dimensions[7].height = 420
    for column in "ABCDEFGH":
        summary.column_dimensions[column].width = 16
    for cell in ("A3", "A4"):
        summary[cell].font = Font(bold=True)
        summary[cell].fill = PatternFill("solid", fgColor="DFF3EE")

    detail = workbook.create_sheet("Structured Response")
    detail.sheet_view.showGridLines = False
    detail.append(["Type", "Content"])
    for cell in detail[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="176B5D")
    for kind, value in _content_blocks(content):
        if kind != "page_break":
            detail.append([kind.replace("_", " ").title(), value])
    detail.freeze_panes = "A2"
    detail.auto_filter.ref = detail.dimensions
    detail.column_dimensions["A"].width = 18
    detail.column_dimensions["B"].width = 110
    for row in detail.iter_rows(min_row=2):
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
    workbook.save(path)


def _content_blocks(content: str) -> list[tuple[str, str]]:
    blocks: list[tuple[str, str]] = []
    paragraph_lines: list[str] = []

    def flush_paragraph() -> None:
        if paragraph_lines:
            blocks.append(("paragraph", " ".join(paragraph_lines).strip()))
            paragraph_lines.clear()

    for raw_line in content.splitlines():
        line = raw_line.strip()
        if not line:
            flush_paragraph()
            continue
        if re.fullmatch(r"[-=_]{3,}", line):
            flush_paragraph()
            continue
        heading = re.sub(r"^\s*#{1,6}\s*", "", line)
        heading = re.sub(r"^\*\*(.+)\*\*$", r"\1", heading)
        numbered_line = bool(re.match(r"^\d+\.\s+\S", line))
        numbered_heading = numbered_line and len(line) <= 80
        plain_heading = heading.lower().rstrip(":") in {
            "analysis report",
            "summary",
            "key observations",
            "recommended next steps",
            "review caveat",
            "draft response prepared",
            "department allegations",
            "taxpayer submission",
            "relevant context retrieved",
            "prayer",
        }
        if line.startswith("#") or numbered_heading or (
            line.startswith("**") and line.endswith("**")
        ) or plain_heading:
            flush_paragraph()
            blocks.append(("heading", heading))
        elif re.match(r"^[-*]\s+", line):
            flush_paragraph()
            blocks.append(("bullet", re.sub(r"^[-*]\s+", "", line)))
        elif numbered_line:
            flush_paragraph()
            blocks.append(("numbered", line))
        else:
            paragraph_lines.append(re.sub(r"\*\*(.+?)\*\*", r"\1", line))
    flush_paragraph()
    return blocks


def _document_title(agent_name: str) -> str:
    return {
        "scn": "Draft Notice Analysis and Response",
        "itr": "Income Tax Return Analysis",
        "bank_statement": "Bank Statement Analysis",
        "financial": "Financial Statement Analysis",
        "ocr": "Document Analysis",
    }.get(agent_name, "CA Analysis Report")


def _slug(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", value.lower()).strip("-")


def _pdf_safe(value: str) -> str:
    return (
        value.replace("₹", "Rs. ")
        .replace("–", "-")
        .replace("—", "-")
        .replace("’", "'")
        .replace("“", '"')
        .replace("”", '"')
    )


def _content_type(output_format: str) -> str:
    return {
        "pdf": "application/pdf",
        "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    }[output_format]
